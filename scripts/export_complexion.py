#!/usr/bin/env python3
"""
Reads every image in the complexion folder and writes a standalone
spreadsheet with two columns -- Image, Complexion -- one row per person,
top to bottom. Meant for copy/pasting into another sheet by hand.

Does NOT open Excel, touch the main workbook, or use either progress file.
Images already present in the output file are skipped, so a run interrupted
by a usage limit can just be re-run. Uncertain/illegible cells are
highlighted yellow.

Usage:
    ./run_complexion_export.sh                 # all images not yet in the export
    ./run_complexion_export.sh --only 98.jpg   # just one image
    ./run_complexion_export.sh --limit 2
"""
import argparse
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from claude_cli import call_claude_vision
from common import SCRIPTS_DIR, add_selection_args, configure_stdout, list_images, load_config, select_images
from prompts import COMPLEXION_PROMPT_TEMPLATE, COMPLEXION_SCHEMA

REVIEW_FILL = PatternFill("solid", fgColor="FFFF00")


def open_output(path: Path):
    if path.is_file():
        wb = load_workbook(path)
        return wb, wb.active
    wb = Workbook()
    ws = wb.active
    ws.title = "Complexion"
    ws.append(["Image", "Complexion"])
    ws["A1"].font = ws["B1"].font = Font(bold=True)
    return wb, ws


def main():
    configure_stdout()
    ap = argparse.ArgumentParser()
    add_selection_args(ap, "config_complexion.json", "98.jpg")
    ap.add_argument("--out", default=str(SCRIPTS_DIR.parent / "complexion_export.xlsx"))
    args = ap.parse_args()

    config = load_config(args.config)
    images_folder = Path(config["images_folder"])
    out_path = Path(args.out)

    if not images_folder.is_dir():
        sys.exit(f"Images folder not found: {images_folder}")

    wb, ws = open_output(out_path)
    already = {row[0] for row in ws.iter_rows(min_row=2, max_col=1, values_only=True) if row[0]}

    todo = select_images(list_images(images_folder), lambda p: p.name in already,
                         args.only, args.limit, images_folder,
                         f"'{{name}}' is already in {out_path}. Delete its rows there to redo it.")

    if not todo:
        print(f"Nothing to do -- every image is already in {out_path}")
        return

    print(f"Exporting {len(todo)} image(s) to {out_path}")
    flagged_total = 0
    consecutive_errors = 0

    for n, img_path in enumerate(todo, 1):
        print(f"[{n}/{len(todo)}] {img_path.name} ...")

        def on_row(i, _img=img_path.name):
            print(f"    reading row {i} of {_img} ...")

        result = call_claude_vision(img_path, images_folder, COMPLEXION_SCHEMA, COMPLEXION_PROMPT_TEMPLATE, on_row=on_row)

        if result["kind"] == "rate_limit":
            print("SUBSCRIPTION LIMIT REACHED.")
            print(f"  {result['message'].strip()}")
            print("Re-run the same command after the reset -- finished images are skipped.")
            return
        if result["kind"] == "error":
            # Keeps going past one bad image, but two in a row usually means
            # something global (e.g. logged out), not the images.
            print(f"ERROR on {img_path.name}: {result['message'][:500]}")
            consecutive_errors += 1
            if consecutive_errors >= 2:
                print("Two images in a row failed -- this is probably not the images themselves "
                      "(e.g. not logged in). Stopping; fix it and re-run the same command.")
                return
            print("Skipping it and moving on.")
            continue
        consecutive_errors = 0

        flagged = 0
        for r in result["rows"]:
            ws.append([img_path.name, r.get("complexion", "")])
            if r.get("status") != "ok":
                flagged += 1
                ws.cell(ws.max_row, 2).fill = REVIEW_FILL
        wb.save(out_path)  # save after every image so nothing is lost if interrupted
        flagged_total += flagged
        print(f"OK ({len(result['rows'])} rows, {flagged} uncertain)")

    print(f"\nDone. Saved: {out_path}  ({flagged_total} uncertain cell(s) highlighted yellow)")


if __name__ == "__main__":
    main()
