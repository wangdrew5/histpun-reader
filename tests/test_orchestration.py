"""The per-image loops of the crime pipeline and the export, with fake
extractors: their deliberately different error handling, and the
write -> save -> mark-done order."""
import io
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import export_complexion  # noqa: E402
from common import load_progress  # noqa: E402
from fake_sheet import FakeBook, FakeSheet  # noqa: E402
from process_reports import process_images  # noqa: E402

HEADERS = ["Name", "Race", "Crime", "Parish", "Home Address", "Status Notes", "File Name"]


def crime_rows(n, tag="x"):
    return [{"row_number": i + 1, "crime": f"{tag}{i}", "parish": "P", "home_address": "A",
             "status": "ok", "notes": ""} for i in range(n)]


def scripted(results):
    """Fake extractor returning the given results in order."""
    calls = []

    def extract(img_path, on_row):
        calls.append(img_path.name)
        return results[len(calls) - 1]
    extract.calls = calls
    return extract


class CrimeLoopTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.progress_path = Path(self.tmp.name) / "progress.json"
        self.sheet = FakeSheet(HEADERS, [{"Name": "existing", "File Name": "1.jpg"}])
        self.book = FakeBook()
        self.progress = {}

    def tearDown(self):
        self.tmp.cleanup()

    def run_loop(self, names, extract, start_row=3):
        with redirect_stdout(io.StringIO()):
            return process_images([Path(n) for n in names], self.sheet, self.book, self.sheet.header_map,
                                  start_row, self.progress, self.progress_path, extract)

    def test_extraction_error_moves_on_to_next_image(self):
        extract = scripted([{"kind": "error", "message": "bad scan"},
                            {"kind": "ok", "rows": crime_rows(2)}])
        summary = self.run_loop(["3.jpg", "5.jpg"], extract)
        self.assertFalse(summary["stopped"])
        saved = load_progress(self.progress_path)
        self.assertEqual(saved["3.jpg"]["status"], "error")
        self.assertEqual(saved["5.jpg"], {**saved["5.jpg"], "status": "done", "excel_rows": [3, 4],
                                          "rows_written": 2})
        self.assertEqual(self.sheet.get(3, "File Name"), "5.jpg")

    def test_pages_append_consecutively(self):
        extract = scripted([{"kind": "ok", "rows": crime_rows(2)}, {"kind": "ok", "rows": crime_rows(3)}])
        self.run_loop(["3.jpg", "5.jpg"], extract)
        saved = load_progress(self.progress_path)
        self.assertEqual(saved["3.jpg"]["excel_rows"], [3, 4])
        self.assertEqual(saved["5.jpg"]["excel_rows"], [5, 7])
        self.assertEqual(self.book.saves, 2)

    def test_occupied_destination_stops_without_writing(self):
        self.sheet.set_record(4, {"Name": "typed mid-run"})
        extract = scripted([{"kind": "ok", "rows": crime_rows(3)}])
        with self.assertRaises(SystemExit):
            self.run_loop(["3.jpg"], extract)
        self.assertEqual(self.sheet.writes, [])
        self.assertEqual(load_progress(self.progress_path)["3.jpg"]["status"], "error")

    def test_not_marked_done_if_workbook_save_fails(self):
        self.book.save = mock.Mock(side_effect=OSError("disk full"))
        extract = scripted([{"kind": "ok", "rows": crime_rows(2)}])
        with self.assertRaises(SystemExit):
            self.run_loop(["3.jpg"], extract)
        self.assertEqual(load_progress(self.progress_path)["3.jpg"]["status"], "in_progress")

    def test_rate_limit_stops_and_records_pending(self):
        extract = scripted([{"kind": "rate_limit", "message": "resets 8pm"}])
        summary = self.run_loop(["3.jpg", "5.jpg"], extract)
        self.assertTrue(summary["stopped"])
        self.assertEqual(extract.calls, ["3.jpg"])
        self.assertEqual(load_progress(self.progress_path)["3.jpg"]["status"], "pending")


class ExportLoopTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        tmp = Path(self.tmp.name)
        self.images = tmp / "images"
        self.images.mkdir()
        for n in range(2, 12, 2):
            (self.images / f"{n}.jpg").write_bytes(b"")
        self.config = tmp / "config.json"
        self.config.write_text('{"images_folder": "%s"}' % self.images)
        self.out = tmp / "export.xlsx"

    def tearDown(self):
        self.tmp.cleanup()

    def run_export(self, results):
        fake = scripted(results)
        argv = ["export_complexion.py", "--config", str(self.config), "--out", str(self.out)]
        with mock.patch.object(export_complexion, "call_claude_vision",
                               side_effect=lambda p, *a, on_row=None, **k: fake(p, on_row)), \
                mock.patch.object(sys, "argv", argv), redirect_stdout(io.StringIO()):
            export_complexion.main()
        return fake.calls

    def exported_images(self):
        from openpyxl import load_workbook
        ws = load_workbook(self.out).active
        return [r[0] for r in ws.iter_rows(min_row=2, max_col=1, values_only=True)]

    def test_stops_after_two_consecutive_errors(self):
        ok = {"kind": "ok", "rows": [{"complexion": "Black", "status": "ok"}]}
        err = {"kind": "error", "message": "nope"}
        calls = self.run_export([err, ok, err, err, ok])
        self.assertEqual(calls, ["2.jpg", "4.jpg", "6.jpg", "8.jpg"])
        self.assertEqual(self.exported_images(), ["4.jpg"])

        # Re-run picks up only what's missing.
        calls = self.run_export([ok, ok, ok, ok])
        self.assertEqual(calls, ["2.jpg", "6.jpg", "8.jpg", "10.jpg"])


if __name__ == "__main__":
    unittest.main()
